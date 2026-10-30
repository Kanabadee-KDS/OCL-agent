"""Skill Transcript set-up (FS-5a–c · M3 · BRD 03 US-M3-01–04).

The skill database (GS / SS / E / C), who sets the skills of each curriculum and with which
rights (the mock's SKP), the skill set of a curriculum per cohort (core + chosen + its own,
approved through the flow `skill_set`), and the course × skill mapping of a term (approved
through `skill_mapping`). A curriculum is a คู่มือหลักสูตร; its students and its offered
sections are those of its programme code in SUPREME, and its courses are those of หมวด 4.
"""

from __future__ import annotations

import io
import re
from typing import Any

from ocl_api.auth import User
from ocl_api.curricula import SCOPE_COLUMNS, visible
from ocl_api.db import Jsonb, connect
from ocl_api.flows import FlowStore, can_act, covers
from ocl_api.mail import Outbox
from ocl_api.skill_defaults import (
    ASSIGN_PERMS,
    ASSIGN_ROLES,
    CATS,
    LEVELS,
    SCOPES,
    SOURCES,
    TOOLS,
    blank_rubric,
)

SET_FLOW, MAP_FLOW = "skill_set", "skill_mapping"
STATUSES = {
    "none": "ยังไม่กำหนด",
    "draft": "ร่าง",
    "sent": "รออนุมัติ",
    "returned": "ส่งคืนแก้ไข",
    "approved": "อนุมัติแล้ว",
}
EDITABLE = ("none", "draft", "returned")
CODE_RE = re.compile(r"^[A-Z]{1,3}\d{1,3}$")
TERM_RE = re.compile(r"^[123]/25\d\d$")
CUR_COLUMNS = (
    "d.id, d.name_th, d.level, d.program_code, d.status AS cur_status, "
    + SCOPE_COLUMNS
    + ", coalesce((SELECT s.data->'courses'->'courses' FROM curriculum_sections s "
    "WHERE s.draft_id = d.id AND s.section = '4'), '[]'::jsonb) AS book_courses"
)


class SkillError(ValueError):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def term_sort(term: str) -> tuple[int, int]:
    n, y = term.split("/")
    return int(y), int(n)


def level_of(score: float | None, levels: list[dict[str, Any]]) -> dict[str, Any] | None:
    """BR-M3-02: the level whose range holds the score (ranges are on 2 decimals)."""
    if score is None:
        return None
    s = round(score, 2)
    for lv in levels:
        if lv["min"] <= s <= lv["max"]:
            return lv
    return levels[-1] if s > levels[-1]["max"] else levels[0]


def pick_set(rows: list[dict[str, Any]], cohort: int | None) -> dict[str, Any] | None:
    """The approved set for students of `cohort` (2-digit entry year), else the latest."""
    rows = sorted([s for s in rows if s["status"] == "approved"], key=lambda s: s["version"])
    if not rows:
        return None
    if cohort is not None:
        fit = [s for s in rows if s["cohort_from"] is None or s["cohort_from"] <= cohort]
        if fit:
            return max(fit, key=lambda s: (s["cohort_from"] or 0, s["version"]))
    return rows[-1]


def _cur(r: dict[str, Any]) -> dict[str, Any]:
    codes = []
    for c in r.get("book_courses") or []:
        code = (c or {}).get("code_th")
        if code and code not in codes:
            codes.append(code)
    return {
        "id": r["id"],
        "name_th": r["name_th"],
        "level": r["level"],
        "program_code": r["program_code"],
        "faculty": r["scope_faculty"],
        "pids": r["scope_pids"] or [],
        "book_courses": codes,
        "scope_faculty": r["scope_faculty"],
        "scope_pids": r["scope_pids"] or [],
    }


def _event(conn, kind: str, ref: str, by: str, action: str, note: str | None = None) -> None:
    conn.execute(
        "INSERT INTO skill_events (kind, ref, by, action, note) VALUES (%s, %s, %s, %s, %s)",
        (kind, ref, by, action, note),
    )


class SkillStore:
    def __init__(self, url: str, schema: str | None = None):
        self.url = url
        self.schema = schema
        self.flows = FlowStore(url, schema)

    def _conn(self):
        return connect(self.url, self.schema)

    # ---------- shared ----------

    def curricula(self, user: User | None = None, *, every: bool = False) -> list[dict]:
        """Curricula the user sees in M3: their ขอบเขตข้อมูล, or those they set skills for
        (`every`: all of them, each marked `seen`)."""
        with self._conn() as conn:
            rows = conn.execute(
                f"SELECT {CUR_COLUMNS} FROM curriculum_drafts d ORDER BY d.name_th, d.id"
            ).fetchall()
            mine = (
                {
                    r["curriculum_id"]
                    for r in conn.execute(
                        "SELECT curriculum_id FROM skill_assign WHERE pid = %s", (user.pid,)
                    )
                }
                if user and user.pid
                else set()
            )
        out = [
            {**_cur(r), "seen": user is None or visible(user, r) or r["id"] in mine} for r in rows
        ]
        return out if every else [c for c in out if c["seen"]]

    def curriculum(self, cid: int, user: User) -> dict[str, Any]:
        cur = next((c for c in self.curricula(user, every=True) if c["id"] == cid), None)
        if cur is None:
            raise SkillError("ไม่พบคู่มือหลักสูตร", 404)
        if not cur["seen"]:
            raise SkillError("หลักสูตรนี้อยู่นอกขอบเขตข้อมูลของบัญชีนี้", 403)
        return cur

    def settings(self) -> dict[str, Any]:
        with self._conn() as conn:
            got = {r["key"]: r["data"] for r in conn.execute("SELECT * FROM skill_settings")}
        return {
            "levels": got.get("levels") or LEVELS,
            "ethics_pass": float(got.get("ethics_pass", 2.5)),
            "frameworks": got.get("frameworks") or [],
        }

    def staff_names(self, pids: list[str] | None = None) -> dict[str, str]:
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT pid, title_short, name_th FROM ext_staff "
                "WHERE %s::text[] IS NULL OR pid = ANY(%s)",
                (pids, pids),
            ).fetchall()
        return {r["pid"]: f"{r['title_short'] or ''}{r['name_th']}" for r in rows}

    def _mail_pids(self, pids: list[str], subject: str, body: str, by: str, ref: str) -> None:
        if not pids:
            return
        with self._conn() as conn:
            to = [
                r["email"]
                for r in conn.execute(
                    "SELECT email FROM ext_staff WHERE pid = ANY(%s) AND email IS NOT NULL",
                    (pids,),
                )
            ]
        if to:
            Outbox(self.url, self.schema).send(to, subject, body, kind="skill", by=by, ref=ref)

    def _mail_roles(
        self, roles: list[str], faculty: str | None, subject: str, body: str, by: str, ref: str
    ) -> None:
        to = self.flows.recipients(roles, faculty)
        if to:
            Outbox(self.url, self.schema).send(to, subject, body, kind="skill", by=by, ref=ref)

    # ---------- FS-5a ฐานข้อมูล Skill (US-M3-01) ----------

    def skills(self, include_inactive: bool = True) -> list[dict[str, Any]]:
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT s.*, (SELECT count(*) FROM skill_sets t WHERE t.skills ? s.code) "
                "AS used_in, d.name_th AS curriculum FROM skills s "
                "LEFT JOIN curriculum_drafts d ON d.id = s.curriculum_id "
                "ORDER BY array_position(ARRAY['GS','SS','E','C'], s.cat), "
                "length(s.code), s.code"
            ).fetchall()
        out = []
        for r in rows:
            if not include_inactive and not r["active"]:
                continue
            out.append({**r, "updated_at": r["updated_at"].isoformat(timespec="seconds")})
        return out

    def groups(self) -> dict[str, str]:
        with self._conn() as conn:
            rows = conn.execute("SELECT id, name FROM role_groups ORDER BY ord").fetchall()
        return {r["id"]: r["name"] for r in rows}

    def db_perms(self) -> list[dict[str, Any]]:
        groups = self.groups()
        with self._conn() as conn:
            rows = conn.execute("SELECT * FROM skill_db_perms ORDER BY id").fetchall()
        names = self.staff_names()
        out = []
        for r in rows:
            kind, _, ref = r["who"].partition(":")
            label = groups.get(ref, ref) if kind == "role" else names.get(ref, ref)
            out.append(
                {
                    "id": r["id"],
                    "who": r["who"],
                    "kind": kind,
                    "label": label,
                    "cats": r["cats"],
                    "unit_scope": r["unit_scope"],
                }
            )
        return out

    def save_db_perm(
        self, pid: int | None, who: str, cats: list[str], unit_scope: str, by: str
    ) -> list[dict[str, Any]]:
        kind, _, ref = who.partition(":")
        if kind not in ("role", "pid") or not ref:
            raise SkillError("ระบุผู้ใช้ (pid:…) หรือกลุ่มสิทธิ์ (role:…)")
        if not cats or set(cats) - set(CATS):
            raise SkillError("เลือกด้านที่เพิ่มได้อย่างน้อย 1 ด้าน")
        if unit_scope not in ("all", "faculty", "program"):
            raise SkillError("ขอบเขตหน่วยงานไม่ถูกต้อง")
        with self._conn() as conn:
            if pid:
                conn.execute(
                    "UPDATE skill_db_perms SET who = %s, cats = %s, unit_scope = %s, "
                    "updated_by = %s, updated_at = now() WHERE id = %s",
                    (who, Jsonb(cats), unit_scope, by, pid),
                )
            else:
                conn.execute(
                    "INSERT INTO skill_db_perms (who, cats, unit_scope, updated_by) "
                    "VALUES (%s, %s, %s, %s)",
                    (who, Jsonb(cats), unit_scope, by),
                )
        return self.db_perms()

    def delete_db_perm(self, pid: int) -> list[dict[str, Any]]:
        with self._conn() as conn:
            conn.execute("DELETE FROM skill_db_perms WHERE id = %s", (pid,))
        return self.db_perms()

    def db_rights(self, user: User) -> dict[str, Any]:
        """Which sides the user may add, and whether also skills of the whole university."""
        if user.sees_all():
            return {"cats": list(CATS), "central": True}
        cats: set[str] = set()
        central = False
        with self._conn() as conn:
            rows = conn.execute("SELECT * FROM skill_db_perms").fetchall()
        for r in rows:
            kind, _, ref = r["who"].partition(":")
            if (kind == "role" and ref in user.roles) or (kind == "pid" and ref == user.pid):
                cats |= set(r["cats"])
                central = central or r["unit_scope"] == "all"
        return {"cats": [c for c in CATS if c in cats], "central": central}

    def _clean_skill(self, body: dict[str, Any]) -> dict[str, Any]:
        code = str(body.get("code") or "").strip().upper()
        if not CODE_RE.match(code):
            raise SkillError("รหัส Skill ต้องเป็นตัวอักษรตามด้วยตัวเลข เช่น SS4")
        cat = body.get("cat")
        if cat not in CATS:
            raise SkillError("ด้านต้องเป็น GS SS E หรือ C")
        name_th = str(body.get("name_th") or "").strip()
        if not name_th:
            raise SkillError("ต้องมีชื่อทักษะ (ไทย)")
        scope = body.get("scope") or "program"
        if scope not in SCOPES:
            raise SkillError("ขอบเขตไม่ถูกต้อง")
        tools = [t for t in body.get("tools") or [] if t in TOOLS]
        sources = [s for s in body.get("sources") or [] if s in SOURCES]
        rubric = [str(x or "") for x in (body.get("rubric") or blank_rubric())][: len(LEVELS)]
        rubric += [""] * (len(LEVELS) - len(rubric))
        return {
            "code": code,
            "cat": cat,
            "name_th": name_th,
            "name_en": (str(body.get("name_en") or "").strip() or None),
            "definition": (str(body.get("definition") or "").strip() or None),
            "scope": scope,
            "tools": tools,
            "sources": sources,
            "rubric": rubric,
        }

    def save_skill(
        self, body: dict[str, Any], user: User, *, create: bool, curriculum_id: int | None = None
    ) -> dict[str, Any]:
        s = self._clean_skill(body)
        rights = self.db_rights(user)
        if curriculum_id is None:
            if s["cat"] not in rights["cats"]:
                raise SkillError(f"บัญชีนี้ไม่มีสิทธิ์เพิ่ม/แก้ Skill ด้าน {s['cat']}", 403)
            if s["scope"] == "central" and not rights["central"]:
                raise SkillError("Skill ส่วนกลาง มศว แก้ได้เฉพาะผู้มีสิทธิ์ระดับมหาวิทยาลัย", 403)
        with self._conn() as conn, conn.transaction():
            old = conn.execute(
                "SELECT * FROM skills WHERE code = %s FOR UPDATE", (s["code"],)
            ).fetchone()
            if create and old:
                raise SkillError(f"มีรหัส {s['code']} แล้ว", 409)
            if not create and not old:
                raise SkillError("ไม่พบ Skill", 404)
            if old and old["scope"] == "central" and not rights["central"]:
                raise SkillError("Skill ส่วนกลาง มศว แก้ได้เฉพาะผู้มีสิทธิ์ระดับมหาวิทยาลัย", 403)
            vals = (
                s["cat"],
                s["name_th"],
                s["name_en"],
                s["definition"],
                s["scope"],
                Jsonb(s["tools"]),
                Jsonb(s["sources"]),
                Jsonb(s["rubric"]),
            )
            if create:
                conn.execute(
                    "INSERT INTO skills (cat, name_th, name_en, definition, scope, tools, "
                    "sources, rubric, code, curriculum_id, created_by, updated_by) "
                    "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                    (*vals, s["code"], curriculum_id, user.email, user.email),
                )
            else:
                conn.execute(
                    "UPDATE skills SET cat = %s, name_th = %s, name_en = %s, definition = %s, "
                    "scope = %s, tools = %s, sources = %s, rubric = %s, updated_by = %s, "
                    "updated_at = now() WHERE code = %s",
                    (*vals, user.email, s["code"]),
                )
            _event(conn, "skill", s["code"], user.email, "create" if create else "update")
        return next(x for x in self.skills() if x["code"] == s["code"])

    def set_active(self, code: str, active: bool, user: User) -> dict[str, Any]:
        with self._conn() as conn:
            old = conn.execute("SELECT * FROM skills WHERE code = %s", (code,)).fetchone()
            if old is None:
                raise SkillError("ไม่พบ Skill", 404)
            if old["cat"] not in self.db_rights(user)["cats"]:
                raise SkillError(f"บัญชีนี้ไม่มีสิทธิ์แก้ Skill ด้าน {old['cat']}", 403)
            conn.execute(
                "UPDATE skills SET active = %s, updated_by = %s, updated_at = now() "
                "WHERE code = %s",
                (active, user.email, code),
            )
            _event(conn, "skill", code, user.email, "enable" if active else "disable")
        return next(x for x in self.skills() if x["code"] == code)

    def save_setting(self, key: str, data: Any, by: str) -> dict[str, Any]:
        if key == "levels":
            rows = list(data or [])
            if len(rows) != len(LEVELS):
                raise SkillError(f"เกณฑ์ระดับต้องมี {len(LEVELS)} ระดับ")
            prev = -1.0
            clean = []
            for i, r in enumerate(rows):
                lo, hi = float(r.get("min")), float(r.get("max"))
                if not (0 <= lo <= hi <= 5) or lo <= prev:
                    raise SkillError(f"ช่วงคะแนนระดับ {i + 1} ต้องอยู่ใน 0–5 และต่อจากระดับก่อน")
                prev = hi
                clean.append(
                    {
                        "level": i + 1,
                        "min": lo,
                        "max": hi,
                        "en": str(r.get("en") or LEVELS[i]["en"]),
                        "th": str(r.get("th") or LEVELS[i]["th"]),
                    }
                )
            data = clean
        elif key == "ethics_pass":
            data = float(data)
            if not 0 <= data <= 5:
                raise SkillError("เกณฑ์ผ่าน Ethics ต้องอยู่ใน 0–5")
        elif key == "frameworks":
            data = [
                {k: str(r.get(k) or "") for k in ("level", "cohorts", "frame", "status")}
                for r in data or []
            ]
        else:
            raise SkillError("ไม่มีการตั้งค่านี้", 404)
        with self._conn() as conn:
            conn.execute(
                "INSERT INTO skill_settings (key, data, updated_by) VALUES (%s, %s, %s) "
                "ON CONFLICT (key) DO UPDATE SET data = EXCLUDED.data, "
                "updated_by = EXCLUDED.updated_by, updated_at = now()",
                (key, Jsonb(data), by),
            )
        return self.settings()

    IMPORT_HEAD = [
        "รหัส",
        "ด้าน",
        "ชื่อทักษะ (ไทย)",
        "ชื่อทักษะ (อังกฤษ)",
        "นิยามเชิงปฏิบัติการ",
        "ขอบเขต (ส่วนกลาง/หลักสูตร)",
    ]

    def template_xlsx(self) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Skill"
        ws.append(self.IMPORT_HEAD)
        ws.append(["SS4", "SS", "[ชื่อทักษะ]", "[English name]", "[นิยาม]", "หลักสูตร"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def import_xlsx(self, data: bytes, user: User) -> dict[str, Any]:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 -- any unreadable file is the user's error
            raise SkillError("อ่านไฟล์ .xlsx ไม่ได้") from e
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        if not rows or [str(c or "").strip() for c in rows[0][:2]] != self.IMPORT_HEAD[:2]:
            raise SkillError("หัวตารางไม่ตรงแม่แบบ (ดาวน์โหลดแม่แบบก่อน)")
        known = {s["code"] for s in self.skills()}
        added = updated = 0
        errors = []
        for n, row in enumerate(rows[1:], start=2):
            cells = [str(c).strip() if c is not None else "" for c in row[:6]] + [""] * 6
            if not any(cells[:6]):
                continue
            body = {
                "code": cells[0],
                "cat": cells[1].upper(),
                "name_th": cells[2],
                "name_en": cells[3],
                "definition": cells[4],
                "scope": "central" if cells[5].startswith("ส่วนกลาง") else "program",
                "tools": list(TOOLS),
                "sources": ["course"],
            }
            try:
                create = body["code"].upper() not in known
                self.save_skill(body, user, create=create)
                known.add(body["code"].upper())
                added, updated = added + create, updated + (not create)
            except SkillError as e:
                errors.append({"row": n, "error": str(e)})
        return {"added": added, "updated": updated, "errors": errors}

    # ---------- FS-5b ผู้รับผิดชอบ (US-M3-02) ----------

    def assigns(self, cid: int | None = None, names: bool = True) -> dict[int, list[dict]]:
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT * FROM skill_assign WHERE %s::bigint IS NULL OR curriculum_id = %s "
                "ORDER BY curriculum_id, updated_at",
                (cid, cid),
            ).fetchall()
        names = self.staff_names([r["pid"] for r in rows]) if names and rows else {}
        out: dict[int, list[dict[str, Any]]] = {}
        for r in rows:
            out.setdefault(r["curriculum_id"], []).append(
                {
                    "pid": r["pid"],
                    "name": names.get(r["pid"], r["pid"]),
                    "role": r["role"],
                    "perms": r["perms"],
                    "scope": r["scope"],
                    "courses": r["courses"],
                }
            )
        return out

    def rights(self, user: User, cur: dict[str, Any]) -> dict[str, Any]:
        """What the user may do on a curriculum's skills (US-M3-02's four rights)."""
        staff = user.sees_all()
        mine = next(
            (
                a
                for a in self.assigns(cur["id"], names=False).get(cur["id"], [])
                if user.pid and a["pid"] == user.pid
            ),
            None,
        )
        perms = set(mine["perms"]) if mine else set()
        faculty = "faculty" in user.scope_kinds() and covers(user.units, cur["faculty"])
        return {
            "staff": staff,
            "assigned": mine is not None,
            "add": staff or "add" in perms,
            "map": staff or "map" in perms,
            "send": staff or "send" in perms,
            "appr": staff or "appr" in perms,
            "faculty": faculty,
            "courses": mine["courses"] if mine and mine["scope"] == "courses" else None,
        }

    def assign_overview(self, user: User) -> dict[str, Any]:
        curs = self.curricula(user)
        got = self.assigns()
        items = []
        for c in curs:
            a = got.get(c["id"], [])
            items.append(
                {
                    "id": c["id"],
                    "name_th": c["name_th"],
                    "faculty": c["faculty"],
                    "level": c["level"],
                    "assign": a,
                    "status": "ok" if a else "none",
                    "courses": self.courses_of(c),
                }
            )
        people = {p["pid"] for c in items for p in c["assign"]}
        return {
            "items": items,
            "kpi": {
                "total": len(items),
                "assigned": sum(1 for i in items if i["assign"]),
                "none": sum(1 for i in items if not i["assign"]),
                "people": len(people),
            },
            "roles": ASSIGN_ROLES,
            "perms": ASSIGN_PERMS,
        }

    def courses_of(self, cur: dict[str, Any]) -> list[str]:
        """Courses of หมวด 4 and of the programme's sections in any term."""
        got = {r["course_code"] for r in self._sections_of(cur)}
        return list(dict.fromkeys([*cur["book_courses"], *sorted(got)]))

    def set_assign(self, cid: int, rows: list[dict[str, Any]], user: User) -> list[dict]:
        cur = self.curriculum(cid, user)
        pids = [r.get("pid") for r in rows]
        if len(set(pids)) != len(pids):
            raise SkillError("มีผู้รับผิดชอบซ้ำ")
        known = self.staff_names(pids) if pids else {}
        offered = set(self.courses_of(cur))
        clean = []
        for r in rows:
            if r.get("pid") not in known:
                raise SkillError(f"ไม่พบบุคลากร {r.get('pid')} ในข้อมูล HURIS")
            if r.get("role") not in ASSIGN_ROLES:
                raise SkillError("บทบาทไม่ถูกต้อง")
            perms = [p for p in ASSIGN_PERMS if p in (r.get("perms") or [])]
            scope = "courses" if r.get("scope") == "courses" else "all"
            courses = (
                [c for c in r.get("courses") or [] if c in offered] if scope == "courses" else []
            )
            if scope == "courses" and not courses:
                raise SkillError(f"{known[r['pid']]} เลือก 'เฉพาะรายวิชา' แต่ยังไม่ได้เลือกรายวิชา")
            clean.append((r["pid"], r["role"], perms, scope, courses))
        with self._conn() as conn, conn.transaction():
            before = {
                x["pid"]
                for x in conn.execute(
                    "SELECT pid FROM skill_assign WHERE curriculum_id = %s", (cid,)
                )
            }
            conn.execute("DELETE FROM skill_assign WHERE curriculum_id = %s", (cid,))
            for pid, role, perms, scope, courses in clean:
                conn.execute(
                    "INSERT INTO skill_assign (curriculum_id, pid, role, perms, scope, courses, "
                    "updated_by) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                    (cid, pid, role, Jsonb(perms), scope, Jsonb(courses), user.email),
                )
            _event(conn, "assign", str(cid), user.email, "assign", ", ".join(known.values()))
        new = [p for p, *_ in clean if p not in before]
        self._mail_pids(
            new,
            f"ได้รับมอบหมายเป็นผู้กำหนด Skill · {cur['name_th']}",
            "ท่านได้รับสิทธิ์กำหนด Skill และ Skill Mapping ของหลักสูตรนี้ในระบบ",
            user.email,
            f"skill-assign:{cid}",
        )
        return self.assigns(cid).get(cid, [])

    # ---------- FS-5b ชุด Skill ของหลักสูตร (US-M3-03) ----------

    def _sets(self, conn, cid: int) -> list[dict[str, Any]]:
        return conn.execute(
            "SELECT * FROM skill_sets WHERE curriculum_id = %s ORDER BY version", (cid,)
        ).fetchall()

    def core(self) -> list[str]:
        """BR-M3-03: the skills of the whole university, in every set, not removable."""
        return [s["code"] for s in self.skills(False) if s["scope"] == "central"]

    def approved_set(self, cid: int, cohort: int | None = None) -> dict[str, Any] | None:
        with self._conn() as conn:
            return pick_set(self._sets(conn, cid), cohort)

    def sets_overview(self, user: User) -> dict[str, Any]:
        curs = self.curricula(user)
        core = set(self.core())
        cats = {s["code"]: s["cat"] for s in self.skills()}
        with self._conn() as conn:
            latest = {
                r["curriculum_id"]: r
                for r in conn.execute(
                    "SELECT DISTINCT ON (curriculum_id) * FROM skill_sets "
                    "ORDER BY curriculum_id, version DESC"
                )
            }
        items = []
        for c in curs:
            s = latest.get(c["id"])
            codes = s["skills"] if s else []
            items.append(
                {
                    "id": c["id"],
                    "name_th": c["name_th"],
                    "faculty": c["faculty"],
                    "status": s["status"] if s else "none",
                    "version": s["version"] if s else None,
                    "cohort_from": s["cohort_from"] if s else None,
                    "counts": {k: sum(1 for x in codes if cats.get(x) == k) for k in CATS},
                    "own": sum(1 for x in codes if x not in core),
                    "only_core": bool(s) and set(codes) <= core,
                }
            )
        st = [i["status"] for i in items]
        return {
            "items": items,
            "statuses": STATUSES,
            "kpi": {
                "approved": st.count("approved"),
                "pending": st.count("draft") + st.count("sent") + st.count("returned"),
                "none": st.count("none"),
                "only_core": sum(1 for i in items if i["only_core"]),
            },
        }

    def skill_set(self, cid: int, user: User) -> dict[str, Any]:
        cur = self.curriculum(cid, user)
        with self._conn() as conn:
            sets = self._sets(conn, cid)
            events = conn.execute(
                "SELECT ts, by, action, note FROM skill_events WHERE kind = 'set' AND ref = %s "
                "ORDER BY id",
                (str(cid),),
            ).fetchall()
        s = sets[-1] if sets else None
        version, steps, _ = self.flows.steps(SET_FLOW, s["flow_version"] if s else None)
        active = self.skills(False)
        core = [x["code"] for x in active if x["scope"] == "central"]
        status = s["status"] if s else "none"
        r = self.rights(user, cur)
        pool = [x for x in active if x["scope"] == "central" or x["curriculum_id"] in (None, cid)]
        return {
            "curriculum": cur,
            "set": {
                "version": s["version"] if s else 1,
                "status": status,
                "status_th": STATUSES[status],
                "skills": s["skills"] if s else core,
                "cohort_from": s["cohort_from"] if s else None,
                "step": s["step"] if s else 0,
            },
            "versions": [
                {
                    "version": x["version"],
                    "status": x["status"],
                    "cohort_from": x["cohort_from"],
                    "skills": len(x["skills"]),
                    "approved_at": x["approved_at"] and x["approved_at"].isoformat(),
                }
                for x in sets
            ],
            "core": core,
            "pool": pool,
            "steps": steps,
            "events": [{**e, "ts": e["ts"].isoformat()} for e in events],
            "rights": r,
            "can_edit": status in EDITABLE and r["add"],
            "can_submit": status in EDITABLE and r["send"],
            "can_certify": status == "sent" and self._may_approve(user, cur, r, steps),
            "can_revise": status == "approved" and r["add"],
            "flow_version": version,
        }

    def _may_approve(self, user: User, cur: dict[str, Any], r: dict[str, Any], steps) -> bool:
        """The faculty of the curriculum (the flow's last step), the curriculum's own approver
        (the right อนุมัติ, as the mock gives the chair) or central."""
        return r["appr"] or (r["faculty"] and can_act(user.roles, steps[-1]))

    def save_set(
        self, cid: int, skills: list[str], cohort_from: int | None, user: User
    ) -> dict[str, Any]:
        d = self.skill_set(cid, user)
        if not d["can_edit"]:
            raise SkillError("บัญชีนี้แก้ชุด Skill นี้ไม่ได้ (ต้องมีสิทธิ์เพิ่ม Skill และชุดยังเป็นร่าง)", 403)
        allowed = {x["code"] for x in d["pool"]}
        bad = [x for x in skills if x not in allowed]
        if bad:
            raise SkillError(f"ไม่มี Skill {', '.join(bad)} ในฐานข้อมูลที่ใช้ได้")
        codes = list(dict.fromkeys([*d["core"], *skills]))  # the core is always in (BR-M3-03)
        if cohort_from is not None and not 0 < cohort_from < 100:
            raise SkillError("รุ่นนิสิตเป็นเลข 2 หลักของปีที่เข้าศึกษา เช่น 67")
        with self._conn() as conn, conn.transaction():
            if d["set"]["status"] == "none":
                conn.execute(
                    "INSERT INTO skill_sets (curriculum_id, version, cohort_from, skills, "
                    "flow_version, updated_by) VALUES (%s, 1, %s, %s, %s, %s)",
                    (cid, cohort_from, Jsonb(codes), d["flow_version"], user.email),
                )
            else:
                conn.execute(
                    "UPDATE skill_sets SET skills = %s, cohort_from = %s, status = 'draft', "
                    "updated_by = %s, updated_at = now() WHERE curriculum_id = %s AND version = %s",
                    (Jsonb(codes), cohort_from, user.email, cid, d["set"]["version"]),
                )
            _event(conn, "set", str(cid), user.email, "save", f"ฉบับ {d['set']['version']}")
        return self.skill_set(cid, user)

    def add_own_skill(self, cid: int, body: dict[str, Any], user: User) -> dict[str, Any]:
        """เพิ่มเอง: a new skill of this curriculum (scope หลักสูตร), straight into its set."""
        d = self.skill_set(cid, user)
        if not d["can_edit"]:
            raise SkillError("บัญชีนี้เพิ่ม Skill ให้หลักสูตรนี้ไม่ได้", 403)
        if body.get("cat") == "GS":
            raise SkillError("ทักษะทั่วไป (GS) เป็นของส่วนกลางเท่านั้น")
        s = self.save_skill({**body, "scope": "program"}, user, create=True, curriculum_id=cid)
        return self.save_set(cid, [*d["set"]["skills"], s["code"]], d["set"]["cohort_from"], user)

    def submit_set(self, cid: int, user: User) -> dict[str, Any]:
        d = self.skill_set(cid, user)
        if not d["can_submit"]:
            raise SkillError("บัญชีนี้ส่งชุด Skill นี้ไม่ได้", 403)
        if d["set"]["status"] == "none":
            self.save_set(cid, d["set"]["skills"], None, user)
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "UPDATE skill_sets SET status = 'sent', step = 1, updated_by = %s, "
                "updated_at = now() WHERE curriculum_id = %s AND version = %s",
                (user.email, cid, d["set"]["version"]),
            )
            _event(conn, "set", str(cid), user.email, "submit", f"ฉบับ {d['set']['version']}")
        cur = d["curriculum"]
        self._mail_roles(
            d["steps"][-1]["roles"],
            cur["faculty"],
            f"ชุด Skill รออนุมัติ · {cur['name_th']}",
            "มีชุด Skill ของหลักสูตรส่งมาให้อนุมัติ",
            user.email,
            f"skill-set:{cid}",
        )
        return self.skill_set(cid, user)

    def certify_set(self, cid: int, ok: bool, note: str, user: User) -> dict[str, Any]:
        d = self.skill_set(cid, user)
        if not d["can_certify"]:
            raise SkillError("บัญชีนี้อนุมัติชุด Skill นี้ไม่ได้", 403)
        if not ok and not note.strip():
            raise SkillError("การส่งคืนต้องระบุเหตุผล")
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "UPDATE skill_sets SET status = %s, step = %s, approved_at = CASE WHEN %s "
                "THEN now() END, updated_by = %s, updated_at = now() "
                "WHERE curriculum_id = %s AND version = %s",
                (
                    "approved" if ok else "returned",
                    2 if ok else 0,
                    ok,
                    user.email,
                    cid,
                    d["set"]["version"],
                ),
            )
            _event(conn, "set", str(cid), user.email, "approve" if ok else "return", note or None)
        senders = [a["pid"] for a in self.assigns(cid).get(cid, []) if "send" in a["perms"]]
        self._mail_pids(
            senders,
            f"ชุด Skill {'อนุมัติแล้ว' if ok else 'ถูกส่งคืนแก้ไข'} · {d['curriculum']['name_th']}",
            note or "-",
            user.email,
            f"skill-set:{cid}",
        )
        return self.skill_set(cid, user)

    def revise_set(self, cid: int, cohort_from: int | None, user: User) -> dict[str, Any]:
        """ขอแก้ไข (สร้างฉบับใหม่): a draft copy of the approved set for a later cohort."""
        d = self.skill_set(cid, user)
        if not d["can_revise"]:
            raise SkillError("สร้างฉบับใหม่ได้เมื่อชุดปัจจุบันอนุมัติแล้ว และมีสิทธิ์เพิ่ม Skill", 403)
        v = d["set"]["version"] + 1
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "INSERT INTO skill_sets (curriculum_id, version, cohort_from, skills, "
                "flow_version, updated_by) VALUES (%s, %s, %s, %s, %s, %s)",
                (cid, v, cohort_from, Jsonb(d["set"]["skills"]), d["flow_version"], user.email),
            )
            _event(conn, "set", str(cid), user.email, "revise", f"ฉบับ {v}")
        return self.skill_set(cid, user)

    # ---------- FS-5c Skill Mapping (US-M3-04) ----------

    def terms(self) -> list[str]:
        with self._conn() as conn:
            got = [r["term"] for r in conn.execute("SELECT DISTINCT term FROM ext_sections")]
        return sorted(got, key=term_sort, reverse=True)

    def _sections_of(self, cur: dict[str, Any], term: str | None = None) -> list[dict]:
        """Sections of the curriculum: its programme code, or a course of its หมวด 4."""
        with self._conn() as conn:
            return conn.execute(
                "SELECT * FROM ext_sections WHERE (%s::text IS NULL OR term = %s) "
                "AND (program_code = %s OR course_code = ANY(%s)) ORDER BY course_code, section",
                (term, term, cur["program_code"] or "-", cur["book_courses"]),
            ).fetchall()

    def mapping(self, term: str, cid: int, user: User) -> dict[str, Any]:
        if not TERM_RE.match(term):
            raise SkillError("ภาคการศึกษาต้องอยู่ในรูป ภาค/ปี เช่น 2/2569")
        cur = self.curriculum(cid, user)
        sset = self.approved_set(cid)
        secs = self._sections_of(cur, term)
        offered = {s["course_code"]: s for s in secs}
        codes = list(dict.fromkeys([*cur["book_courses"], *offered]))
        with self._conn() as conn:
            central = {
                r["code_th"]: r
                for r in conn.execute(
                    "SELECT c.code_th, c.module, v.name_th FROM courses c "
                    "LEFT JOIN course_versions v ON v.code_th = c.code_th AND v.version = c.latest "
                    "WHERE c.code_th = ANY(%s)",
                    (codes,),
                )
            }
            m = conn.execute(
                "SELECT * FROM skill_maps WHERE term = %s AND curriculum_id = %s", (term, cid)
            ).fetchone()
            events = conn.execute(
                "SELECT ts, by, action, note FROM skill_events WHERE kind = 'map' AND ref = %s "
                "ORDER BY id",
                (f"{term}|{cid}",),
            ).fetchall()
        version, steps, _ = self.flows.steps(MAP_FLOW, m["flow_version"] if m else None)
        cells = (m["cells"] if m else {}) or {}
        skills = sset["skills"] if sset else []
        courses = [
            {
                "code": c,
                "name_th": (central.get(c) or {}).get("name_th")
                or (offered.get(c) or {}).get("course_name")
                or c,
                "module": (central.get(c) or {}).get("module"),
                "offered": c in offered,
                "skills": [x for x in cells.get(c, []) if x in skills],
            }
            for c in codes
        ]
        live = [c for c in courses if c["offered"]]
        status = m["status"] if m else "none"
        r = self.rights(user, cur)
        return {
            "term": term,
            "curriculum": cur,
            "set": sset
            and {"version": sset["version"], "cohort_from": sset["cohort_from"], "skills": skills},
            "skills": [s for s in self.skills() if s["code"] in skills],
            "courses": courses,
            "status": status,
            "status_th": STATUSES[status],
            "steps": steps,
            "events": [{**e, "ts": e["ts"].isoformat()} for e in events],
            "kpi": {
                "courses": len(live),
                "mapped": sum(1 for c in live if c["skills"]),
                "points": sum(len(c["skills"]) for c in live),
                "gaps": [s for s in skills if not any(s in c["skills"] for c in live)],
            },
            "rights": r,
            "can_edit": bool(sset) and status in EDITABLE and r["map"],
            "can_submit": bool(sset) and status in EDITABLE and r["send"],
            "can_certify": status == "sent" and self._may_approve(user, cur, r, steps),
            "flow_version": version,
        }

    def save_mapping(
        self, term: str, cid: int, cells: dict[str, list[str]], user: User
    ) -> dict[str, Any]:
        d = self.mapping(term, cid, user)
        if not d["can_edit"]:
            raise SkillError("บัญชีนี้แก้ Skill Mapping นี้ไม่ได้", 403)
        live = {c["code"] for c in d["courses"] if c["offered"]}
        old = {c["code"]: c["skills"] for c in d["courses"]}
        mine = d["rights"]["courses"]
        skills = set(d["set"]["skills"])
        clean: dict[str, list[str]] = {}
        for code, xs in cells.items():
            want = sorted({x for x in xs if x in skills}, key=list(d["set"]["skills"]).index)
            if code not in live:
                if want:
                    raise SkillError(f"รายวิชา {code} ไม่เปิดสอนในภาคนี้ กำหนด Skill ไม่ได้")
                continue
            if mine is not None and code not in mine and want != old.get(code, []):
                raise SkillError(f"บัญชีนี้กำหนดได้เฉพาะรายวิชาที่ได้รับมอบหมาย ({code})", 403)
            if want:
                clean[code] = want
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "INSERT INTO skill_maps (term, curriculum_id, cells, flow_version, updated_by) "
                "VALUES (%s, %s, %s, %s, %s) ON CONFLICT (term, curriculum_id) DO UPDATE SET "
                "cells = EXCLUDED.cells, status = 'draft', updated_by = EXCLUDED.updated_by, "
                "updated_at = now()",
                (term, cid, Jsonb(clean), d["flow_version"], user.email),
            )
            _event(conn, "map", f"{term}|{cid}", user.email, "save")
        return self.mapping(term, cid, user)

    def copy_mapping(self, term: str, cid: int, user: User) -> dict[str, Any]:
        """คัดลอกจากภาคก่อน: the latest earlier term's mapping, for courses offered now."""
        d = self.mapping(term, cid, user)
        if not d["can_edit"]:
            raise SkillError("บัญชีนี้แก้ Skill Mapping นี้ไม่ได้", 403)
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT term, cells FROM skill_maps WHERE curriculum_id = %s", (cid,)
            ).fetchall()
        earlier = sorted(
            [r for r in rows if term_sort(r["term"]) < term_sort(term)],
            key=lambda r: term_sort(r["term"]),
        )
        if not earlier:
            raise SkillError("ไม่มี Skill Mapping ของภาคก่อนให้คัดลอก", 404)
        src = earlier[-1]
        live = {c["code"] for c in d["courses"] if c["offered"]}
        cells = {c["code"]: c["skills"] for c in d["courses"] if c["skills"]}
        skipped = []
        for code, xs in (src["cells"] or {}).items():
            if code in live:
                cells[code] = sorted(set(cells.get(code, [])) | set(xs))
            else:
                skipped.append(code)
        out = self.save_mapping(term, cid, cells, user)
        return {**out, "copied_from": src["term"], "skipped": skipped}

    def submit_mapping(self, term: str, cid: int, user: User) -> dict[str, Any]:
        d = self.mapping(term, cid, user)
        if not d["can_submit"]:
            raise SkillError("บัญชีนี้ส่ง Skill Mapping นี้ไม่ได้", 403)
        if not d["kpi"]["points"]:
            raise SkillError("ยังไม่ได้กำหนด Skill ให้รายวิชาใด")
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "INSERT INTO skill_maps (term, curriculum_id, status, step, flow_version, "
                "updated_by) VALUES (%s, %s, 'sent', 1, %s, %s) "
                "ON CONFLICT (term, curriculum_id) DO UPDATE SET status = 'sent', step = 1, "
                "updated_by = EXCLUDED.updated_by, updated_at = now()",
                (term, cid, d["flow_version"], user.email),
            )
            _event(conn, "map", f"{term}|{cid}", user.email, "submit")
        cur = d["curriculum"]
        self._mail_roles(
            d["steps"][-1]["roles"],
            cur["faculty"],
            f"Skill Mapping รออนุมัติ · {cur['name_th']} ภาค {term}",
            "มี Skill Mapping ส่งมาให้อนุมัติ",
            user.email,
            f"skill-map:{term}:{cid}",
        )
        return self.mapping(term, cid, user)

    def certify_mapping(
        self, term: str, cid: int, ok: bool, note: str, user: User
    ) -> dict[str, Any]:
        d = self.mapping(term, cid, user)
        if not d["can_certify"]:
            raise SkillError("บัญชีนี้อนุมัติ Skill Mapping นี้ไม่ได้", 403)
        if not ok and not note.strip():
            raise SkillError("การส่งคืนต้องระบุเหตุผล")
        with self._conn() as conn, conn.transaction():
            conn.execute(
                "UPDATE skill_maps SET status = %s, step = %s, approved_at = CASE WHEN %s "
                "THEN now() END, updated_by = %s, updated_at = now() "
                "WHERE term = %s AND curriculum_id = %s",
                ("approved" if ok else "returned", 2 if ok else 0, ok, user.email, term, cid),
            )
            _event(
                conn,
                "map",
                f"{term}|{cid}",
                user.email,
                "approve" if ok else "return",
                note or None,
            )
        senders = [a["pid"] for a in self.assigns(cid).get(cid, []) if "send" in a["perms"]]
        self._mail_pids(
            senders,
            f"Skill Mapping {'อนุมัติแล้ว' if ok else 'ถูกส่งคืนแก้ไข'} · "
            f"{d['curriculum']['name_th']} ภาค {term}",
            note or "-",
            user.email,
            f"skill-map:{term}:{cid}",
        )
        return self.mapping(term, cid, user)

    def approved_maps(self, term: str) -> dict[str, dict[str, Any]]:
        """BR-M3-06: course -> skills to assess in the term, from approved mappings only
        (a course mapped by two curricula is assessed on the union), with the curricula."""
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT m.curriculum_id, m.cells, d.name_th FROM skill_maps m "
                "JOIN curriculum_drafts d ON d.id = m.curriculum_id "
                "WHERE m.term = %s AND m.status = 'approved'",
                (term,),
            ).fetchall()
        out: dict[str, dict[str, Any]] = {}
        for r in rows:
            for code, xs in (r["cells"] or {}).items():
                e = out.setdefault(code, {"skills": [], "curricula": []})
                e["skills"] = list(dict.fromkeys([*e["skills"], *xs]))
                e["curricula"].append({"id": r["curriculum_id"], "name_th": r["name_th"]})
        return out
