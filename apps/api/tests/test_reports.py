"""Reports R01–R14, the home dashboard and R11 (FS-7 · US-M5-01–04 · US-G-01)."""

from __future__ import annotations

import io

from tests.test_skills import (  # noqa: F401 -- `env` is the fixture
    ADVISOR,
    CENTRAL,
    CHAIR,
    SLUG,
    STUDENT,
    TEACHER,
    _score_all,
    _set_up,
    env,
)


def test_catalog_follows_the_menus(env) -> None:  # noqa: F811
    c = env["c"]
    codes = [r["code"] for r in c.get("/api/reports", headers=CENTRAL).json()]
    assert codes == [f"R{n:02d}" for n in range(1, 15)]
    teacher = {r["code"] for r in c.get("/api/reports", headers=TEACHER).json()}
    assert teacher == {"R07", "R08", "R11", "R13", "R14"}, "r0 r7 r8 (and the public R11)"
    assert c.get("/api/reports", headers=STUDENT).status_code == 403
    assert c.get("/api/reports/R01", headers=TEACHER).status_code == 403
    assert c.get("/api/reports/R99", headers=CENTRAL).status_code == 404
    pub = c.get("/api/public/stats").json()
    assert pub["code"] == "R11" and pub["kpis"][0]["value"] >= 1, "no sign-in (US-M5-01)"


def test_reports_carry_the_data_of_their_module(env) -> None:  # noqa: F811
    from openpyxl import load_workbook

    c = env["c"]
    _set_up(env)
    _score_all(c, "ทวพ131", "1", {})
    r01 = c.get(f"/api/reports/R01?term={SLUG.replace('-', '/')}", headers=CENTRAL).json()
    assert r01["kpis"][0]["label"] == "เปิดสอน" and r01["remind"] and len(r01["tables"]) == 2
    r02 = c.get("/api/reports/R02", headers=CENTRAL).json()
    assert r02["kpis"][0]["value"] == 4 and r02["kpis"][1]["value"] == 1
    assert c.post("/api/reports/R02/remind", headers=CENTRAL).json()["sent"] == 3
    assert c.post("/api/reports/R02/remind", headers=CENTRAL).json()["sent"] == 0, "once a day"
    r03 = c.get("/api/reports/R03", headers=CENTRAL).json()
    assert any(row["name"] == "วท.บ. วัสดุพอลิเมอร์" for row in r03["tables"][0]["rows"])
    # skills: the confirmed section counts; the advisor sees their own students only
    r04 = c.get("/api/reports/R04", headers=CENTRAL).json()
    ss1 = next(r for r in r04["tables"][0]["rows"] if r["code"] == "SS1")
    assert ss1["avg"] == 4.0 and ss1["n"] == 30 and ss1["courses"] == 1
    r05 = c.get("/api/reports/R05?by=entry_year", headers=CENTRAL).json()
    assert r05["tables"][0]["rows"][0]["group"] == "ทั้งมหาวิทยาลัย"
    assert c.get("/api/reports/R06", headers=CENTRAL).json()["tables"][1]["rows"][0]["SS"] == 4.0
    r07 = c.get("/api/reports/R07?section=ทวพ131|1", headers=CENTRAL).json()
    assert len(r07["tables"][0]["rows"]) == 30 and r07["tables"][0]["rows"][0]["SS1"] == 4.0
    r08 = c.get("/api/reports/R08", headers=ADVISOR).json()
    assert len(r08["filters"][0]["options"]) < 40, "นิสิตในความดูแล"
    # R12: twelve sheets in one Excel file, named with the code and the date (US-M5-04)
    x = c.get("/api/reports/R12.xlsx?all=1", headers=CENTRAL)
    assert x.status_code == 200 and "R12-" in x.headers["content-disposition"]
    assert len(load_workbook(io.BytesIO(x.content)).sheetnames) == 12
    c.post("/api/courses", json={"code_th": "ทวพ999", "name_th": "วิชากลาง", "code_en": "X", "name_en": "X",
           "transcript_th": "x", "transcript_en": "X", "course_type": "วิชาแกน",
           "credit_text": "3(2-2-5)", "status": "ใช้งาน"}, headers=CENTRAL)  # fmt: skip
    r14 = c.get("/api/reports/R14", headers=CENTRAL).json()
    assert r14["kpis"][0]["value"] == len(r14["tables"][0]["rows"]) >= 1
    r13 = c.get(f"/api/reports/R13?curriculum={env['cid']}", headers=CHAIR).json()
    assert r13["filters"][0]["value"] == str(env["cid"])


def test_r09_counts_checks_and_home(env) -> None:  # noqa: F811
    c, grads = env["c"], env["grads"]
    _set_up(env)
    for code in ("ทวพ111", "ทวพ121", "ทวพ131"):
        _score_all(c, code, "1", {})
    no = c.post("/api/skills/issue", json={"sids": [grads[0]]}, headers=CENTRAL).json()["issued"][0]
    for ip in ("1.1.1.1", "1.1.1.1", "2.2.2.2"):
        c.get(f"/api/public/transcripts/{no}", headers={"x-forwarded-for": ip})
    r09 = c.get("/api/reports/R09", headers=CENTRAL).json()
    assert [k["value"] for k in r09["kpis"]] == [1, 3, 2], "issued · checks · distinct checkers"
    assert r09["tables"][1]["rows"][0]["n"] == 3
    row = r09["tables"][0]["rows"][0]
    assert row["status"] == "ออกแล้ว" and len(row["issued_at"]) == 16, "Thai label, Thai time"
    h = c.get("/api/home", headers=CENTRAL).json()
    assert {x["label"] for x in h["cards"]} >= {"หลักสูตรในระบบ", "ประเมิน Skill ครบ"}
    assert h["skills"]["GS"] == 4.0
    mine = c.get("/api/home", headers=TEACHER).json()
    assert mine["cards"] == [] and all("ประเมิน Skill" in t["text"] for t in mine["tasks"])
